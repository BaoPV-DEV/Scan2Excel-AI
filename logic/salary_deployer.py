import os
import re
import openpyxl
from openpyxl.utils import get_column_letter
from logic.salary_config_parser import get_salary_config

def apply_salary_formulas_from_json(excel_files_dir, template_key, month, year, log_callback, progress_callback=None):
    """
    Áp dụng công thức lương từ JSON config cho các file Excel trong thư mục.
    Tự động mapping đến D:\Linh_Salary_Tool\{YYYY}\{MM}\ để tìm file source.
    Nếu file/sheet source chưa có → để trống ô + log cảnh báo.
    """
    if not os.path.exists(excel_files_dir):
        return False, f"❌ Thư mục không tồn tại: {excel_files_dir}"
    
    try:
        config = get_salary_config()
    except Exception as e:
        return False, f"❌ Lỗi đọc config JSON: {str(e)}"
    
    # Validate template
    is_valid, errors = config.validate_template(template_key)
    if not is_valid:
        return False, f"❌ Template config lỗi:\n" + "\n".join(errors)
    
    log_callback(f"📊 Bắt đầu triển khai công thức lương")
    log_callback(f"📋 Template: {template_key}")
    log_callback(f"📂 Thư mục: {excel_files_dir}")
    log_callback(f"📅 Tháng/Năm: {month}/{year}")
    log_callback(f"🔧 Đọc config từ: {config.config_path}\n")
    
    # Tính toán các giá trị thời gian động
    yyyy = int(year)
    mm = int(month)
    yyyy_str = str(yyyy)
    mm_str = f"{mm:02d}"
    mm_int = mm
    
    # Tính tháng trước
    if mm_int == 1:
        prev_month_int = 12
        prev_year_int = yyyy - 1
    else:
        prev_month_int = mm_int - 1
        prev_year_int = yyyy
    
    prev_month = f"{prev_month_int:02d}"
    prev_year = str(prev_year_int)
    yy_short = yyyy_str[-2:]
    prev_yy_short = prev_year[-2:]
    
    # Tìm tất cả file Excel trong thư mục (lọc theo template name)
    template_filename = os.path.basename(template_key)  # VD: bep_an_va_cong_vu_template.xlsx
    excel_files = []
    for root, dirs, files in os.walk(excel_files_dir):
        for file in files:
            if file.endswith('.xlsx') and template_filename in file:
                excel_files.append(os.path.join(root, file))
    
    if not excel_files:
        return False, f"❌ Không tìm thấy file '{template_filename}' nào trong thư mục."
    
    log_callback(f"🔍 Tìm thấy {len(excel_files)} file Excel để xử lý\n")
    
    template_config = config.get_template_config(template_key)
    processed = 0
    warnings = []
    
    for file_idx, excel_file in enumerate(excel_files):
        try:
            log_callback(f"\n📄 [{file_idx + 1}/{len(excel_files)}] Xử lý: {os.path.basename(excel_file)}")
            
            wb = openpyxl.load_workbook(excel_file)
            file_modified = False
            
            # Lặp qua tất cả sheets trong template config
            for sheet_pattern, sheet_cfg in template_config.get("sheets", {}).items():
                # Resolve dynamic sheet name
                resolved_sheet_name = sheet_pattern.replace("{mm}", mm_str)
                
                # Tìm sheet trong workbook (case-insensitive)
                ws = None
                actual_sheet_name = None
                for name in wb.sheetnames:
                    if name.lower() == resolved_sheet_name.lower():
                        ws = wb[name]
                        actual_sheet_name = name
                        break
                
                if not ws:
                    warning_msg = f"   ⚠️ Sheet '{resolved_sheet_name}' không tìm thấy → Bỏ qua"
                    log_callback(warning_msg)
                    warnings.append(warning_msg)
                    continue
                
                # Mở khóa sheet
                try:
                    ws.protection.sheet = False
                except Exception:
                    pass
                
                start_row = sheet_cfg.get("start_row", 9)
                row_step = sheet_cfg.get("row_step", 2)
                
                # Tính số lượng nhân viên (từ dữ liệu trong ô C)
                num_employees = 0
                for row in range(start_row, start_row + 500, row_step):
                    cell_value = ws.cell(row=row, column=3).value
                    if cell_value and str(cell_value).strip():
                        num_employees += 1
                    else:
                        break
                
                if num_employees == 0:
                    log_callback(f"   ℹ️ Sheet '{actual_sheet_name}': Không có dữ liệu nhân viên")
                    continue
                
                # Xác định tỷ lệ cho Bếp/Công vụ
                bep_cong_vu_rate = "100%"
                file_lower = os.path.basename(excel_file).lower()
                if "bếp" in file_lower or "bep" in file_lower:
                    bep_cong_vu_rate = "125%"
                elif "công vụ" in file_lower or "cong vu" in file_lower:
                    bep_cong_vu_rate = "115%"
                
                log_callback(f"   📋 Sheet: {actual_sheet_name}")
                log_callback(f"   👥 Nhân viên: {num_employees} dòng (row {start_row})")
                
                # Lặp qua từng cột có config
                columns_config = sheet_cfg.get("columns", {})
                for col_letter, col_cfg in columns_config.items():
                    parsed = config.parse_column_config(col_cfg)
                    col_num = ord(col_letter.upper()) - ord('A') + 1
                    
                    if parsed["type"] == "formula":
                        # Công thức tĩnh
                        formula_template = parsed["formula"]
                        log_callback(f"      📌 Cột {col_letter}: Công thức")
                        
                        for idx in range(num_employees):
                            current_row = start_row + (idx * row_step)
                            formula = formula_template.format(
                                row=current_row,
                                row_minus_1=current_row - 1,
                                row_plus_1=current_row + 1,
                                bep_cong_vu_rate=bep_cong_vu_rate
                            )
                            try:
                                ws.cell(row=current_row, column=col_num).value = formula
                            except Exception as e:
                                log_callback(f"         ❌ Lỗi: {str(e)}")
                        file_modified = True
                    
                    elif parsed["type"] == "source":
                        # Linking đến file ngoài
                        source_file = parsed["file"]
                        source_sheet = parsed["sheet"]
                        source_range = parsed["range"]
                        col_index = parsed.get("column_index", 1)
                        
                        # Resolve đường dẫn file source
                        source_path = config.resolve_file_path(
                            source_file, yyyy, mm, log_callback
                        )
                        
                        if source_path is None:
                            # File không tồn tại → để trống + warning
                            warning = f"      ⚠️ Cột {col_letter}: File '{source_file}' không tìm thấy → Để trống"
                            log_callback(warning)
                            warnings.append(warning)
                            
                            # Xóa data (để trống)
                            for idx in range(num_employees):
                                current_row = start_row + (idx * row_step)
                                ws.cell(row=current_row, column=col_num).value = None
                            file_modified = True
                        else:
                            # File tồn tại → tạo formula linking
                            log_callback(f"      ✅ Cột {col_letter}: Link từ {os.path.basename(source_path)}!{source_sheet}")
                            
                            for idx in range(num_employees):
                                current_row = start_row + (idx * row_step)
                                # Build Excel external link formula
                                formula = f"=IFERROR(VLOOKUP(C{current_row},'{source_path}[{os.path.basename(source_path)}]{source_sheet}'!{source_range},{col_index},0),\"\")"
                                try:
                                    ws.cell(row=current_row, column=col_num).value = formula
                                except Exception as e:
                                    log_callback(f"         ❌ Lỗi: {str(e)}")
                            file_modified = True
            
            # Lưu file nếu có thay đổi
            if file_modified:
                wb.save(excel_file)
                processed += 1
                log_callback(f"   ✅ Đã lưu file")
            
            wb.close()
            
            if progress_callback:
                progress_callback(int(((file_idx + 1) / len(excel_files)) * 100))
            
        except Exception as e:
            log_callback(f"   ❌ Lỗi xử lý file: {str(e)}")
            import traceback
            log_callback(f"   {traceback.format_exc()}")
    
    if progress_callback:
        progress_callback(100)
    
    result_msg = f"\n✅ Hoàn tất! Đã triển khai cho {processed}/{len(excel_files)} file."
    if warnings:
        result_msg += f"\n\n⚠️ Có {len(warnings)} cảnh báo - Kiểm tra log ở trên."
    
    return True, result_msg
