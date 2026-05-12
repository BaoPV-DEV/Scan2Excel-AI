import pandas as pd
import openpyxl
import os
import shutil
import copy
import numpy as np
import pythoncom
import win32com.client as win32

def process_excel(input_path, template_path, output_dir, mm, yyyy, log_callback, progress_callback=None):
    """
    Hàm xử lý tách file Excel danh sách nhân viên thành từng file riêng cho mỗi tổ/bộ phận.
    """
    mm_yyyy = f"{mm}/{yyyy}"
    
    log_callback(f"=== Bắt đầu xử lý ===")
    log_callback(f"Tháng/Năm báo cáo: {mm_yyyy}")
    log_callback(f"Đang đọc file đầu vào: {os.path.basename(input_path)}...")

    try:
        # ============================================================
        # BƯỚC 1: Đọc file Excel đầu vào - sheet "Danh sách"
        # ============================================================
        df = pd.read_excel(input_path, sheet_name="Danh sách", header=None)

        log_callback(f"Tổng số hàng đọc được: {len(df)}")
        log_callback(f"Tổng số cột: {df.shape[1]}")

        # Định nghĩa vị trí các cột dựa trên cấu trúc file mẫu (C, D, E)
        COL_TO = 2     # Cột C - Tổ/bộ phận
        COL_ID = 3     # Cột D - Mã số nhân viên
        COL_NAME = 4   # Cột E - Họ và Tên

        if df.shape[1] <= COL_NAME:
            raise ValueError(f"File input không đủ cột (cần ít nhất cột E, hiện chỉ có {df.shape[1]} cột).")

        # ============================================================
        # BƯỚC 2: Làm sạch và chuẩn bị dữ liệu
        # ============================================================
        # Điền giá trị Tổ/bộ phận xuống cho các ô trống (ffill)
        df[COL_TO] = df[COL_TO].replace(r'^\s*$', np.nan, regex=True)
        df[COL_TO] = df[COL_TO].ffill()

        # Chuẩn hóa kiểu dữ liệu (về chuỗi và xóa khoảng trắng)
        df[COL_TO] = df[COL_TO].astype(str).str.strip()
        df[COL_NAME] = df[COL_NAME].astype(str).str.strip()
        df[COL_ID] = df[COL_ID].astype(str).str.replace(r'\.0$', '', regex=True).str.strip()

        # Thay thế các giá trị rỗng thành NA để dễ dàng drop
        df[COL_TO] = df[COL_TO].replace(['nan', 'None', ''], pd.NA)
        df[COL_ID] = df[COL_ID].replace(['nan', 'None', ''], pd.NA)
        df[COL_NAME] = df[COL_NAME].replace(['nan', 'None', ''], pd.NA)

        # Loại bỏ các hàng không hợp lệ (thiếu thông tin quan trọng)
        df_valid = df.dropna(subset=[COL_TO, COL_ID, COL_NAME])

        # Loại bỏ các hàng tiêu đề (nếu có lẫn trong dữ liệu)
        header_mask = (
            df_valid[COL_TO].str.lower().str.contains('tổ/bộ phận|bộ phận', na=False, regex=True) &
            df_valid[COL_ID].str.lower().str.contains('mã', na=False)
        )
        df_valid = df_valid[~header_mask]
        df_valid = df_valid[~df_valid[COL_TO].str.match(r'^\d+$', na=False)]

        # Lưu lại tên hiển thị gốc cho Tổ/bộ phận (tránh biến thành chữ thường hoàn toàn)
        df_valid = df_valid.copy()
        df_valid['_to_normalized'] = df_valid[COL_TO].str.lower()

        display_name_map = {}
        for _, row in df_valid.iterrows():
            norm = row['_to_normalized']
            if norm not in display_name_map:
                display_name_map[norm] = row[COL_TO]

        df_valid[COL_TO] = df_valid['_to_normalized'].map(display_name_map)

        if df_valid.empty:
            raise ValueError("Không tìm thấy dữ liệu hợp lệ sau khi lọc. Vui lòng kiểm tra lại file input.")

        # ============================================================
        # BƯỚC 3: Nhóm dữ liệu theo Tổ/bộ phận
        # ============================================================
        unique_groups = df_valid[COL_TO].unique()
        log_callback(f"Tìm thấy {len(unique_groups)} Tổ/bộ phận duy nhất:")
        for g in unique_groups:
            count = len(df_valid[df_valid[COL_TO] == g])
            log_callback(f"  - {g}: {count} nhân viên")

        # ============================================================
        # BƯỚC 4: Kiểm tra file mẫu (Template)
        # ============================================================
        log_callback(f"\nĐang kiểm tra file mẫu: {os.path.basename(template_path)}...")
        wb_check = openpyxl.load_workbook(template_path, read_only=True)
        if "Bang TH nop" not in wb_check.sheetnames:
            raise ValueError(f"Sheet 'Bang TH nop' không tồn tại trong file template. "
                             f"Các sheet có: {', '.join(wb_check.sheetnames)}")
        wb_check.close()
        log_callback("File mẫu OK - Đã tìm thấy sheet 'Bang TH nop'.")

        # ============================================================
        # BƯỚC 5: Xử lý từng tổ và ghi vào Excel (sử dụng win32com để giữ format)
        # ============================================================
        processed_count = 0
        processed_groups = set()

        pythoncom.CoInitialize() # Khởi tạo môi trường COM cho đa luồng
        excel_app = None
        try:
            # Khởi tạo ứng dụng Excel ẩn danh
            excel_app = win32.DispatchEx('Excel.Application')
            excel_app.Visible = False
            excel_app.DisplayAlerts = False
            excel_app.ScreenUpdating = False
            excel_app.AskToUpdateLinks = False
            excel_app.Interactive = False

            total_groups = len(unique_groups)
            for i, to_name in enumerate(unique_groups):
                if to_name in processed_groups:
                    continue
                processed_groups.add(to_name)

                # Cập nhật tiến độ UI
                if progress_callback:
                    progress_callback(int((i / total_groups) * 100))

                group_data = df_valid[df_valid[COL_TO] == to_name].reset_index(drop=True)

                # Chuẩn hóa tên tổ để hiển thị và đặt tên file
                display_to_name = str(to_name)
                if display_to_name.isdigit():
                    display_to_name = f"Tổ {display_to_name}"
                
                log_callback(f"\n--- Đang tạo file cho: {display_to_name} ({len(group_data)} nhân viên) ---")

                # Tạo tên file an toàn (xóa ký tự đặc biệt)
                safe_to_name = display_to_name.replace('/', '_').replace('\\', '_')
                output_filename = f"{safe_to_name} - {mm}.xlsx"
                output_filepath = os.path.join(output_dir, output_filename)

                # Kiểm tra và ghi đè file cũ nếu cần
                if os.path.exists(output_filepath):
                    try:
                        os.remove(output_filepath)
                    except PermissionError:
                        log_callback(f"  ❌ Lỗi: Không thể ghi đè '{output_filename}'. Vui lòng đóng file nếu đang mở.")
                        continue

                # Copy file mẫu sang file đích
                shutil.copy2(template_path, output_filepath)

                # Mở file đích bằng win32com để bắt đầu điền dữ liệu
                wb = excel_app.Workbooks.Open(os.path.abspath(output_filepath), UpdateLinks=0)
                ws = wb.Sheets("Bang TH nop")

                # Cập nhật tiêu đề bảng (Tháng và Tên Tổ)
                header_text = f"BẢNG TỔNG HỢP GIÂY THÁNG {mm_yyyy}- {str(to_name).upper()}"
                header_col = 1
                for c in range(1, 11):
                    val = ws.Cells(2, c).Value
                    if val and isinstance(val, str) and "BẢNG" in val.upper():
                        header_col = c
                        break
                ws.Cells(2, header_col).Value = header_text

                # Điền danh sách nhân viên (Bắt đầu từ hàng 9, mỗi người cách 2 hàng do file mẫu)
                START_ROW = 9
                ROW_STEP = 2

                for idx, (_, row_data) in enumerate(group_data.iterrows()):
                    current_row = START_ROW + (idx * ROW_STEP)
                    ws.Cells(current_row, 1).Value = idx + 1
                    ws.Cells(current_row, 2).Value = row_data[COL_NAME]
                    ws.Cells(current_row, 3).Value = row_data[COL_ID]

                # Lưu và đóng file
                wb.Save()
                wb.Close(False)
                processed_count += 1
                log_callback(f"  ✅ Thành công: {output_filename}")

            if progress_callback:
                progress_callback(100)
                
        finally:
            # Đảm bảo thoát ứng dụng Excel để không chạy ngầm
            if excel_app:
                excel_app.Quit()
            pythoncom.CoUninitialize()

        # ============================================================
        # BƯỚC 6: Tổng kết kết quả
        # ============================================================
        log_callback(f"\n{'='*50}")
        log_callback(f"=== HOÀN THÀNH QUÁ TRÌNH ===")
        log_callback(f"Tổng số file đã tạo: {processed_count}/{len(unique_groups)}")
        log_callback(f"Thư mục lưu trữ: {output_dir}")
        log_callback(f"{'='*50}")
        
        return True, f"Tách file hoàn tất!\nĐã tạo {processed_count} file thành công."

    except FileNotFoundError as e:
        log_callback(f"❌ Lỗi: Không tìm thấy file - {str(e)}")
        return False, f"Không tìm thấy file:\n{str(e)}"
    except ValueError as e:
        log_callback(f"❌ Lỗi dữ liệu: {str(e)}")
        return False, f"Lỗi dữ liệu:\n{str(e)}"
    except PermissionError as e:
        log_callback(f"❌ Lỗi quyền truy cập: {str(e)}")
        return False, f"Lỗi quyền truy cập (File đang mở):\n{str(e)}"
    except Exception as e:
        log_callback(f"❌ Lỗi không xác định: {str(e)}")
        return False, f"Đã xảy ra lỗi:\n{str(e)}"

